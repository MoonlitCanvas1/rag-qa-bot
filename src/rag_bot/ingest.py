"""Load documents from disk, split them into overlapping chunks, and store them."""

import hashlib
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

if TYPE_CHECKING:
    from langchain_chroma import Chroma

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset({".txt", ".md", ".pdf", ".docx"})
DEFAULT_BATCH_SIZE = 500


class IngestError(Exception):
    """Raised when nothing can be ingested from the given path."""


@dataclass(frozen=True)
class IngestStats:
    files: int
    skipped: int
    chunks: int


# --------------------------------------------------------------------------- discovery


def discover_files(root: Path) -> list[Path]:
    """Return supported files under ``root`` (or ``root`` itself if it is a file), sorted."""
    if root.is_file():
        return [root] if root.suffix.lower() in SUPPORTED_EXTENSIONS else []
    return sorted(
        p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def source_label(path: Path, base: Path) -> str:
    """Stable, human-readable source name: path relative to ``base`` when possible."""
    try:
        return path.resolve().relative_to(base.resolve()).as_posix()
    except ValueError:
        return path.name


# --------------------------------------------------------------------------- loading


def _read_text(path: Path, source: str) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [Document(page_content=text, metadata={"source": source})] if text else []


def _read_pdf(path: Path, source: str) -> list[Document]:
    reader = PdfReader(path)
    docs: list[Document] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(
                Document(page_content=text, metadata={"source": source, "page": page_number})
            )
    return docs


def _read_docx(path: Path, source: str) -> list[Document]:
    doc = DocxDocument(str(path))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts).strip()
    return [Document(page_content=text, metadata={"source": source})] if text else []


def load_file(path: Path, source: str | None = None) -> list[Document]:
    """Load one file into LangChain documents (one per page for PDFs)."""
    suffix = path.suffix.lower()
    label = source or path.name
    if suffix in {".txt", ".md"}:
        return _read_text(path, label)
    if suffix == ".pdf":
        return _read_pdf(path, label)
    if suffix == ".docx":
        return _read_docx(path, label)
    raise IngestError(f"Unsupported file type: {path.suffix!r} ({path})")


# --------------------------------------------------------------------------- chunking


def _chunk_id(source: str, index: int, text: str) -> str:
    digest = hashlib.sha1(f"{source}:{index}:{text}".encode()).hexdigest()
    return digest[:32]


def chunk_documents(docs: list[Document], chunk_size: int, chunk_overlap: int) -> list[Document]:
    """Split documents into overlapping chunks.

    Each chunk keeps its parent's metadata and gains ``chunk_index`` (per source,
    across pages) and a deterministic ``chunk_id``.
    """
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    counters: dict[str, int] = {}
    chunks: list[Document] = []
    for doc in docs:
        source = str(doc.metadata.get("source", "unknown"))
        for text in splitter.split_text(doc.page_content):
            index = counters.get(source, 0)
            counters[source] = index + 1
            metadata = {
                **doc.metadata,
                "chunk_index": index,
                "chunk_id": _chunk_id(source, index, text),
            }
            chunks.append(Document(page_content=text, metadata=metadata))
    return chunks


# --------------------------------------------------------------------------- storage


def _replace_source(store: "Chroma", source: str, chunks: list[Document], batch_size: int) -> None:
    """Delete existing chunks for ``source`` then add the new ones (idempotent re-ingest)."""
    existing = store.get(where={"source": source}).get("ids", [])
    if existing:
        logger.debug("Removing %d stale chunks for %s", len(existing), source)
        store.delete(ids=existing)
    for start in range(0, len(chunks), batch_size):
        batch = chunks[start : start + batch_size]
        store.add_documents(batch, ids=[str(d.metadata["chunk_id"]) for d in batch])


def ingest_path(
    path: Path,
    store: "Chroma",
    *,
    chunk_size: int,
    chunk_overlap: int,
    base_dir: Path | None = None,
    batch_size: int = DEFAULT_BATCH_SIZE,
) -> IngestStats:
    """Ingest a file or directory into ``store``."""
    if not path.exists():
        raise IngestError(f"Path does not exist: {path}")

    files = discover_files(path)
    if not files:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise IngestError(f"No supported files ({supported}) found at {path}")

    base = base_dir or (path if path.is_dir() else path.parent)
    ingested = skipped = total_chunks = 0

    for file in files:
        source = source_label(file, base)
        try:
            docs = load_file(file, source)
        except Exception as exc:  # one corrupt file must not abort the whole run
            logger.warning("Skipping %s: %s", source, exc)
            skipped += 1
            continue

        chunks = chunk_documents(docs, chunk_size, chunk_overlap)
        if not chunks:
            logger.warning("Skipping %s: no extractable text", source)
            skipped += 1
            continue

        _replace_source(store, source, chunks, batch_size)
        ingested += 1
        total_chunks += len(chunks)
        logger.info("Ingested %s (%d chunks)", source, len(chunks))

    return IngestStats(files=ingested, skipped=skipped, chunks=total_chunks)
