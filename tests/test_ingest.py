from pathlib import Path
from typing import Any

import pytest
from docx import Document as DocxDocument
from langchain_core.documents import Document

from rag_bot.ingest import (
    IngestError,
    chunk_documents,
    discover_files,
    ingest_path,
    load_file,
    source_label,
)


class FakeStore:
    """Minimal in-memory stand-in for the Chroma methods ingest uses."""

    def __init__(self) -> None:
        self.docs: dict[str, Document] = {}

    def get(self, where: dict[str, Any]) -> dict[str, list[str]]:
        (key, value), = where.items()
        return {"ids": [i for i, d in self.docs.items() if d.metadata.get(key) == value]}

    def delete(self, ids: list[str]) -> None:
        for i in ids:
            self.docs.pop(i, None)

    def add_documents(self, documents: list[Document], ids: list[str]) -> None:
        self.docs.update(zip(ids, documents, strict=True))


def numbered_words(n: int) -> str:
    return " ".join(f"w{i:03d}" for i in range(n))


# ------------------------------------------------------------------ chunking


def test_chunks_respect_size_limit() -> None:
    doc = Document(page_content=numbered_words(200), metadata={"source": "a.txt"})
    chunks = chunk_documents([doc], chunk_size=100, chunk_overlap=20)
    assert len(chunks) > 1
    assert all(len(c.page_content) <= 100 for c in chunks)


def test_chunks_overlap() -> None:
    doc = Document(page_content=numbered_words(200), metadata={"source": "a.txt"})
    chunks = chunk_documents([doc], chunk_size=100, chunk_overlap=30)
    for prev, nxt in zip(chunks, chunks[1:], strict=False):
        first_word = nxt.page_content.split()[0]
        assert first_word in prev.page_content.split()[-8:]


def test_chunks_cover_all_words() -> None:
    doc = Document(page_content=numbered_words(120), metadata={"source": "a.txt"})
    chunks = chunk_documents([doc], chunk_size=100, chunk_overlap=20)
    seen = {w for c in chunks for w in c.page_content.split()}
    assert seen == set(doc.page_content.split())


def test_chunk_metadata_and_ids() -> None:
    docs = [
        Document(page_content=numbered_words(60), metadata={"source": "a.pdf", "page": 1}),
        Document(page_content=numbered_words(60), metadata={"source": "a.pdf", "page": 2}),
    ]
    chunks = chunk_documents(docs, chunk_size=100, chunk_overlap=20)
    assert [c.metadata["chunk_index"] for c in chunks] == list(range(len(chunks)))
    assert {c.metadata["page"] for c in chunks} == {1, 2}
    ids = [c.metadata["chunk_id"] for c in chunks]
    assert len(set(ids)) == len(ids)

    again = chunk_documents(docs, chunk_size=100, chunk_overlap=20)
    assert ids == [c.metadata["chunk_id"] for c in again]


def test_overlap_must_be_smaller_than_size() -> None:
    with pytest.raises(ValueError):
        chunk_documents([], chunk_size=50, chunk_overlap=50)


# ------------------------------------------------------------------ loading


def test_load_txt_and_md(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("hello world", encoding="utf-8")
    (tmp_path / "b.md").write_text("# Title\n\nbody", encoding="utf-8")
    assert load_file(tmp_path / "a.txt")[0].page_content == "hello world"
    assert "body" in load_file(tmp_path / "b.md")[0].page_content


def test_load_docx(tmp_path: Path) -> None:
    path = tmp_path / "doc.docx"
    d = DocxDocument()
    d.add_paragraph("First paragraph")
    d.add_paragraph("Second paragraph")
    d.save(str(path))
    text = load_file(path, "doc.docx")[0].page_content
    assert "First paragraph" in text and "Second paragraph" in text


def test_load_empty_file_yields_nothing(tmp_path: Path) -> None:
    (tmp_path / "empty.txt").write_text("   \n", encoding="utf-8")
    assert load_file(tmp_path / "empty.txt") == []


def test_load_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "x.csv"
    path.write_text("a,b")
    with pytest.raises(IngestError):
        load_file(path)


def test_discover_files_filters_and_recurses(tmp_path: Path) -> None:
    (tmp_path / "sub").mkdir()
    for name in ("a.txt", "b.MD", "sub/c.docx", "skip.csv", ".gitkeep"):
        (tmp_path / name).write_text("x")
    found = [p.relative_to(tmp_path).as_posix() for p in discover_files(tmp_path)]
    assert found == ["a.txt", "b.MD", "sub/c.docx"]


def test_source_label_relative_and_fallback(tmp_path: Path) -> None:
    (tmp_path / "sub").mkdir()
    inside = tmp_path / "sub" / "a.txt"
    assert source_label(inside, tmp_path) == "sub/a.txt"
    assert source_label(Path("/elsewhere/b.txt"), tmp_path) == "b.txt"


# ------------------------------------------------------------------ ingest_path


def test_ingest_path_stores_chunks_and_skips_bad_files(tmp_path: Path) -> None:
    (tmp_path / "good.txt").write_text(numbered_words(100), encoding="utf-8")
    (tmp_path / "blank.txt").write_text("", encoding="utf-8")
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")

    store = FakeStore()
    stats = ingest_path(tmp_path, store, chunk_size=100, chunk_overlap=20)  # type: ignore[arg-type]

    assert stats.files == 1
    assert stats.skipped == 2
    assert stats.chunks == len(store.docs) > 1
    assert {d.metadata["source"] for d in store.docs.values()} == {"good.txt"}


def test_reingest_replaces_stale_chunks(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"
    store = FakeStore()

    path.write_text(numbered_words(200), encoding="utf-8")
    first = ingest_path(tmp_path, store, chunk_size=100, chunk_overlap=20)  # type: ignore[arg-type]

    path.write_text(numbered_words(30), encoding="utf-8")
    second = ingest_path(tmp_path, store, chunk_size=100, chunk_overlap=20)  # type: ignore[arg-type]

    assert second.chunks < first.chunks
    assert len(store.docs) == second.chunks


def test_ingest_path_errors(tmp_path: Path) -> None:
    store = FakeStore()
    with pytest.raises(IngestError):
        ingest_path(tmp_path / "missing", store, chunk_size=100, chunk_overlap=20)  # type: ignore[arg-type]
    with pytest.raises(IngestError):
        ingest_path(tmp_path, store, chunk_size=100, chunk_overlap=20)  # type: ignore[arg-type]
